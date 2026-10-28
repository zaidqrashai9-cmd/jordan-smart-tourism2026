"""Read the official DOS 2025 bilingual locality workbook into an attributed directory."""
import json,pathlib,hashlib,re,sys
from openpyxl import load_workbook
ROOT=pathlib.Path(__file__).resolve().parents[1]
URL='https://dosweb.dos.gov.jo/DataBank/Population/Population_Estimares/PopulationEstimatesbyLocality.xlsx'
GOV={'العاصمة':'Amman','البلقاء':'Balqa','الزرقاء':'Zarqa','مادبا':'Madaba','إربد':'Irbid','المفرق':'Mafraq','جرش':'Jerash','عجلون':'Ajloun','الكرك':'Karak','الطفيلة':'Tafilah','معان':'Maan','العقبة':'Aqaba'}
def clean(s):return re.sub(r'\s+',' ',str(s or '')).strip()
def build(path):
 w=load_workbook(path,read_only=True,data_only=True);items=[]
 for s in w:
  gov=s.title.strip()
  if gov not in GOV:continue
  district='';district_en='';sub=''
  for i,r in enumerate(s.iter_rows(values_only=True),1):
   if len(r)<8:continue
   if isinstance(r[0],str) and r[0].strip() not in ['المجموع','اللواء']:district=clean(r[0])
   if len(r)>9 and isinstance(r[9],str) and r[9].strip() not in ['Total','District']:district_en=clean(r[9])
   if isinstance(r[1],str) and r[1].strip() not in ['القضاء']:sub=clean(r[1])
   ar=clean(r[2]);en=clean(r[7])
   if not ar or ar in ['التجمع','المجموع'] or en.lower()=='total' or not isinstance(r[5],(int,float)):continue
   key=gov+'|'+ar+'|'+en+'|'+district+'|'+sub+'|'+str(i)
   items.append({'id':hashlib.sha256(key.encode()).hexdigest()[:16],'ar':ar,'en':en,'governorate':gov,'governorate_en':GOV[gov],'district':district,'district_en':district_en,'subdistrict':sub,'type':'locality','aliases_ar':list(dict.fromkeys([ar,ar[:-1]+'ة' if ar.endswith('ه') else ar])),'aliases_en':[en] if en else [],'source_row':i,'source_sheet':s.title,'english_status':'official' if en else 'missing'})
 govs=[{'id':'gov-'+en.lower(),'ar':ar,'en':en,'governorate':ar,'governorate_en':en,'district':'','district_en':'','subdistrict':'','type':'governorate','aliases_ar':[ar]+(['عمان','عمّان'] if en=='Amman' else []),'aliases_en':[en],'english_status':'official'} for ar,en in GOV.items()]
 out={'source':URL,'source_name':'دائرة الإحصاءات العامة الأردنية — تقديرات السكان حسب التجمع 2025','source_year':2025,'retrieved':'2026-10-28','sha256':hashlib.sha256(pathlib.Path(path).read_bytes()).hexdigest(),'completeness':'All numeric locality rows in the 12 governorate worksheets; locality is the source classification, not a claim that every row is a city. Combined localities remain combined.','locality_count':len(items),'governorate_count':12,'missing_english':sum(not x['en'] for x in items),'items':govs+items}
 (ROOT/'monitoring/localities.json').write_text(json.dumps(out,ensure_ascii=False,separators=(',',':')))
 (ROOT/'ios/Resources/localities.json').write_text(json.dumps(out,ensure_ascii=False,separators=(',',':')))
 print(json.dumps({k:v for k,v in out.items() if k!='items'},ensure_ascii=False))
if __name__=='__main__':build(sys.argv[1])
